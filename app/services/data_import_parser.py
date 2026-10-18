"""Shared typed row validation, strict JSON parsing, and Excel workbooks.

Both transports normalize their raw cells into transport-neutral mappings and
call one shared conversion layer, so a field can never gain one meaning through
JSON and another through Excel. Transport-specific code only handles cell syntax
and header or key structure:

- JSON adds strict ISO dates, JSON booleans, duplicate-key detection, and
  root-shape enforcement.
- Excel adds native workbook dates, booleans, numeric cells, eight-digit
  ``YYYYMMDD`` dates, customer-name lookup, and container/formula rejection.

Parsing is synchronous; the worker (Task 9) runs it in a thread.
"""

import json
import re
import zipfile
from collections.abc import Iterator, Mapping, MutableMapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Protocol, cast
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import (  # type: ignore[import-untyped]
    InvalidFileException,
)
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import BusinessError, ErrorCode
from app.models.contact import ContactStatus
from app.models.customer import Customer
from app.models.data_management import (
    DataFieldDefinition,
    DataFieldRole,
    DataFieldType,
)
from app.repositories.data_table import ResolvedDataDefinition
from app.services.customer_name import customer_name_lookup_key

_MAX_ROWS = 100_000
_MAX_TEXT_LENGTH = 4_000
_INT64_MIN = -(2**63)
_INT64_MAX = 2**63 - 1
_MAX_DECIMAL_INTEGER_DIGITS = 22
_MAX_DECIMAL_PLACES = 8
_ISO_DATE_PATTERN = re.compile(r"^[0-9]{4}-[0-9]{2}-[0-9]{2}$")
_INTEGER_PATTERN = re.compile(r"^-?[0-9]+$")
_DECIMAL_PATTERN = re.compile(
    r"^-?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?$"
)
_FIELD_ID_PATTERN = re.compile(r"^[1-9][0-9]{0,18}$")

_XLSX_SIGNATURE = b"PK\x03\x04"
_XLSX_MAX_BYTES = 50 * 1024 * 1024
_WORKBOOK_PART = "xl/workbook.xml"
_WORKBOOK_RELS_PART = "xl/_rels/workbook.xml.rels"
_MACRO_CONTENT_TYPE = b"macroEnabled"
_MACRO_PART = "xl/vbaProject.bin"
_CHARTSHEET_PREFIX = "xl/chartsheets/"
_EXCEL_SHEET_COUNT_MESSAGE = "Excel file must contain exactly one visible worksheet"
_EXCEL_HEADER_MESSAGE = "Excel header row is invalid"
_EXCEL_INVALID_MESSAGE = "Import Excel file is invalid"
_EXCEL_FORMULA_MESSAGE = "Excel formulas are not supported"
_EXCEL_ERROR_VALUE_MESSAGE = "Excel error values are not supported"
_EXCEL_EXTRA_VALUE_MESSAGE = "Excel row contains values outside the configured columns"
_EXCEL_BOOLEAN_TRUE = frozenset({"true", "是", "1"})
_EXCEL_BOOLEAN_FALSE = frozenset({"false", "否", "0"})
_FIXED_JSON_KEYS: dict[DataFieldRole, str] = {
    DataFieldRole.BUSINESS_DATE: "business_date",
    DataFieldRole.UPSTREAM_CUSTOMER: "upstream_customer_id",
    DataFieldRole.DOWNSTREAM_CUSTOMER: "downstream_customer_id",
    DataFieldRole.REVENUE: "revenue",
    DataFieldRole.COST: "cost",
}
_FIXED_JSON_KEY_SET = frozenset(_FIXED_JSON_KEYS.values())


@dataclass(frozen=True)
class CustomerIdentity:
    """The small active-customer projection used during import validation."""

    id: int
    name: str
    status: ContactStatus


@dataclass(frozen=True)
class FieldValue:
    """One custom field value mapped to the typed value-table columns."""

    field_id: int
    text_value: str | None = None
    integer_value: int | None = None
    decimal_value: Decimal | None = None
    date_value: date | None = None
    boolean_value: bool | None = None
    enum_option_id: int | None = None


@dataclass(frozen=True)
class RowIssue:
    """A safe, row-scoped validation error with optional field location."""

    error_code: str
    message: str
    field_id: int | None = None
    column_label: str | None = None
    sheet_name: str | None = None
    row_number: int | None = None


@dataclass(frozen=True)
class ValidatedImportRow:
    """A fully typed row ready for snapshot persistence."""

    business_date: date
    upstream_customer_id: int
    downstream_customer_id: int
    upstream_customer_name_source: str
    downstream_customer_name_source: str
    revenue: Decimal
    cost: Decimal
    custom_values: Mapping[int, FieldValue]

    @property
    def field_values(self) -> tuple[FieldValue, ...]:
        return tuple(self.custom_values.values())

    @property
    def text_values(self) -> dict[int, str]:
        return {
            field_id: value.text_value
            for field_id, value in self.custom_values.items()
            if value.text_value is not None
        }

    @property
    def integer_values(self) -> dict[int, int]:
        return {
            field_id: value.integer_value
            for field_id, value in self.custom_values.items()
            if value.integer_value is not None
        }

    @property
    def decimal_values(self) -> dict[int, Decimal]:
        return {
            field_id: value.decimal_value
            for field_id, value in self.custom_values.items()
            if value.decimal_value is not None
        }

    @property
    def date_values(self) -> dict[int, date]:
        return {
            field_id: value.date_value
            for field_id, value in self.custom_values.items()
            if value.date_value is not None
        }

    @property
    def boolean_values(self) -> dict[int, bool]:
        return {
            field_id: value.boolean_value
            for field_id, value in self.custom_values.items()
            if value.boolean_value is not None
        }

    @property
    def enum_option_ids(self) -> dict[int, int]:
        return {
            field_id: value.enum_option_id
            for field_id, value in self.custom_values.items()
            if value.enum_option_id is not None
        }


@dataclass(frozen=True)
class ParsedImportRow:
    """One source row, either validated or carrying all row-scoped issues."""

    row_number: int
    row: ValidatedImportRow | None = None
    issues: tuple[RowIssue, ...] = ()


class _JsonObject(MutableMapping[str, object]):
    """A decoded object that retains duplicate-key evidence."""

    def __init__(self, pairs: list[tuple[str, object]]) -> None:
        values: dict[str, object] = {}
        duplicate_keys: list[str] = []
        for key, value in pairs:
            if key in values:
                duplicate_keys.append(key)
            values[key] = value
        self._values = values
        self.duplicate_keys = tuple(duplicate_keys)

    def __getitem__(self, key: str) -> object:
        return self._values[key]

    def __setitem__(self, key: str, value: object) -> None:
        self._values[key] = value

    def __delitem__(self, key: str) -> None:
        del self._values[key]

    def __iter__(self) -> Iterator[str]:
        return iter(self._values)

    def __len__(self) -> int:
        return len(self._values)


class _ConversionError(Exception):
    def __init__(self, error_code: str, message: str) -> None:
        super().__init__(message)
        self.error_code = error_code
        self.message = message


class _CellCodec(Protocol):
    """Transport-specific normalization for one raw field value.

    The shared validator owns required checks, field-type dispatch, enum label
    matching, issue construction, and customer scope enforcement. A codec only
    decides how a transport's cell syntax becomes the shared domain value, so
    JSON and Excel can never diverge on validation rules.
    """

    def text(self, value: object) -> str: ...

    def integer(self, value: object) -> int: ...

    def decimal(self, value: object) -> Decimal: ...

    def date(self, value: object) -> date: ...

    def boolean(self, value: object) -> bool: ...

    def enum(
        self,
        value: object,
        field: DataFieldDefinition,
        definition: ResolvedDataDefinition,
    ) -> int: ...

    def customer_id(
        self,
        value: object,
        customers_by_name: Mapping[str, CustomerIdentity],
    ) -> int: ...


class _JsonCellCodec:
    """Strict JSON literals: ISO dates, real booleans, no binary floats."""

    def text(self, value: object) -> str:
        return _coerce_text(value)

    def integer(self, value: object) -> int:
        return _coerce_integer(value)

    def decimal(self, value: object) -> Decimal:
        return _coerce_decimal(value)

    def date(self, value: object) -> date:
        return _coerce_date(value)

    def boolean(self, value: object) -> bool:
        return _coerce_boolean(value)

    def enum(
        self,
        value: object,
        field: DataFieldDefinition,
        definition: ResolvedDataDefinition,
    ) -> int:
        return _coerce_enum(value, field, definition)

    def customer_id(
        self,
        value: object,
        customers_by_name: Mapping[str, CustomerIdentity],
    ) -> int:
        return _coerce_customer_id(value)


_JSON_CELL_CODEC = _JsonCellCodec()


def _invalid_json() -> BusinessError:
    return BusinessError(
        ErrorCode.DATA_IMPORT_INVALID,
        "Import JSON is invalid",
        http_status=422,
    )


def _reject_json_constant(constant: str) -> object:
    raise ValueError(f"Non-standard JSON constant: {constant}")


def _load_json(path: Path) -> object:
    try:
        with path.open(encoding="utf-8") as source:
            value = json.load(
                source,
                parse_float=Decimal,
                parse_constant=_reject_json_constant,
                object_pairs_hook=_JsonObject,
            )
    except (
        json.JSONDecodeError,
        UnicodeDecodeError,
        OSError,
        RecursionError,
        ValueError,
        InvalidOperation,
    ) as error:
        raise _invalid_json() from error
    return cast(object, value)


def _load_json_rows(path: Path) -> list[object]:
    document = _load_json(path)
    if not isinstance(document, Mapping):
        raise _invalid_json()
    if isinstance(document, _JsonObject) and document.duplicate_keys:
        raise _invalid_json()
    if set(document) != {"rows"}:
        raise _invalid_json()
    rows = document["rows"]
    if not isinstance(rows, list) or len(rows) > _MAX_ROWS:
        raise _invalid_json()
    return rows


def json_import_row_count(path: Path) -> int:
    """Return the row count after enforcing the complete JSON envelope."""
    return len(_load_json_rows(path))


def iter_json_rows(
    path: Path,
    definition: ResolvedDataDefinition,
    customer_index: Mapping[str, CustomerIdentity],
    allowed_customer_ids: set[int] | None,
) -> Iterator[ParsedImportRow]:
    """Yield one validated or issue-carrying result per JSON row."""
    customers_by_id = _customer_index_by_id(customer_index)
    for row_number, raw in enumerate(_load_json_rows(path), start=1):
        if not isinstance(raw, Mapping):
            yield ParsedImportRow(
                row_number=row_number,
                issues=(RowIssue("invalid_row", "Row must be a JSON object"),),
            )
            continue

        duplicate_issues = _duplicate_issue_rows(raw, definition, row_number)
        validation = _validate_raw_row(
            raw,
            definition,
            customers_by_id,
            allowed_customer_ids,
        )
        if isinstance(validation, ValidatedImportRow) and not duplicate_issues:
            yield ParsedImportRow(row_number=row_number, row=validation)
            continue

        validation_issues = validation if isinstance(validation, list) else []
        issues = [
            RowIssue(
                error_code=issue.error_code,
                message=issue.message,
                field_id=issue.field_id,
                column_label=issue.column_label,
                sheet_name=issue.sheet_name,
                row_number=row_number,
            )
            for issue in [*duplicate_issues, *validation_issues]
        ]
        yield ParsedImportRow(row_number=row_number, issues=tuple(issues))


async def load_customer_name_index(
    session: AsyncSession,
    visible_customer_ids: set[int] | None = None,
) -> dict[str, CustomerIdentity]:
    """Load active customer names, optionally restricted to visible IDs."""
    if visible_customer_ids is not None and not visible_customer_ids:
        return {}

    statement = select(Customer).where(
        Customer.status == ContactStatus.ACTIVE,
        Customer.deleted_at.is_(None),
    )
    if visible_customer_ids is not None:
        statement = statement.where(Customer.id.in_(visible_customer_ids))
    customers = list((await session.scalars(statement)).all())
    return {
        customer_name_lookup_key(customer.name): CustomerIdentity(
            id=customer.id,
            name=customer.name.strip(),
            status=customer.status,
        )
        for customer in customers
    }


def _customer_index_by_id(
    customer_index: Mapping[str, CustomerIdentity],
) -> dict[int, CustomerIdentity]:
    return {identity.id: identity for identity in customer_index.values()}


def validate_raw_row(
    raw: Mapping[str, object],
    definition: ResolvedDataDefinition,
    customer_index: Mapping[str, CustomerIdentity],
    allowed_customer_ids: set[int] | None,
) -> ValidatedImportRow | list[RowIssue]:
    """Validate one transport-neutral row against a resolved definition."""
    return _validate_raw_row(
        raw,
        definition,
        _customer_index_by_id(customer_index),
        allowed_customer_ids,
    )


def _validate_raw_row(
    raw: Mapping[str, object],
    definition: ResolvedDataDefinition,
    customers_by_id: Mapping[int, CustomerIdentity],
    allowed_customer_ids: set[int] | None,
) -> ValidatedImportRow | list[RowIssue]:
    """Validate one JSON row against an already-built customer-ID index."""
    issues: list[RowIssue] = []
    system_raw: dict[DataFieldRole, object] = {}
    for role, raw_key in _FIXED_JSON_KEYS.items():
        system_raw[role] = raw.get(raw_key)

    structural_issues: list[RowIssue] = []
    allowed_keys = _FIXED_JSON_KEY_SET | {"custom_values"}
    for key in raw:
        if key not in allowed_keys:
            structural_issues.append(
                RowIssue(
                    "unknown_field",
                    "Row contains an unknown field",
                )
            )

    custom_values_raw = raw.get("custom_values", {})
    if custom_values_raw is None:
        custom_values_raw = {}
    if not isinstance(custom_values_raw, Mapping):
        structural_issues.append(
            RowIssue(
                "invalid_custom_values",
                "custom_values must be an object keyed by field ID",
            )
        )
        custom_values_raw = {}

    custom_raw = _json_custom_values(
        custom_values_raw,
        definition,
        structural_issues,
    )
    return _validate_values(
        system_raw,
        custom_raw,
        definition,
        customers_by_id,
        _JSON_CELL_CODEC,
        issues,
        allowed_customer_ids=allowed_customer_ids,
        structural_issues=structural_issues,
    )


def _validate_values(
    system_raw: Mapping[DataFieldRole, object],
    custom_raw: Mapping[int, object],
    definition: ResolvedDataDefinition,
    customers_by_id: Mapping[int, CustomerIdentity],
    codec: _CellCodec,
    issues: list[RowIssue],
    *,
    allowed_customer_ids: set[int] | None,
    customers_by_name: Mapping[str, CustomerIdentity] | None = None,
    structural_issues: Sequence[RowIssue] = (),
) -> ValidatedImportRow | list[RowIssue]:
    """Shared conversion for both transports.

    Required checks, field-type dispatch, enum matching, customer scope, and
    issue construction live here exactly once; only ``codec`` decides how a
    transport spells a value. Customer scope is always enforced here after the
    codec resolves the ID, so JSON and Excel cannot diverge.

    ``structural_issues`` are transport-level key/column findings (unknown JSON
    keys, system fields nested in ``custom_values``) appended in the historic
    order: system field issues, structural issues, then custom field issues.
    """
    system_values: dict[DataFieldRole, object] = {}
    for role in _FIXED_JSON_KEYS:
        field = _required_system_field(definition, role)
        raw_value = system_raw.get(role)
        if raw_value is None:
            issues.append(_required_issue(field))
            continue
        try:
            if role == DataFieldRole.BUSINESS_DATE:
                system_values[role] = codec.date(raw_value)
            elif role in (
                DataFieldRole.UPSTREAM_CUSTOMER,
                DataFieldRole.DOWNSTREAM_CUSTOMER,
            ):
                customer_id = codec.customer_id(
                    raw_value,
                    customers_by_name or {},
                )
                system_values[role] = _resolve_customer(
                    customer_id,
                    customers_by_id,
                    allowed_customer_ids,
                )
            else:
                system_values[role] = codec.decimal(raw_value)
        except _ConversionError as error:
            issues.append(_conversion_issue(field, error))

    issues.extend(structural_issues)
    custom_values = _convert_custom_values(custom_raw, definition, codec, issues)
    if issues:
        return issues

    upstream = cast(
        CustomerIdentity,
        system_values[DataFieldRole.UPSTREAM_CUSTOMER],
    )
    downstream = cast(
        CustomerIdentity,
        system_values[DataFieldRole.DOWNSTREAM_CUSTOMER],
    )
    return ValidatedImportRow(
        business_date=cast(date, system_values[DataFieldRole.BUSINESS_DATE]),
        upstream_customer_id=upstream.id,
        downstream_customer_id=downstream.id,
        upstream_customer_name_source=upstream.name,
        downstream_customer_name_source=downstream.name,
        revenue=cast(Decimal, system_values[DataFieldRole.REVENUE]),
        cost=cast(Decimal, system_values[DataFieldRole.COST]),
        custom_values=custom_values,
    )


def _required_system_field(
    definition: ResolvedDataDefinition,
    role: DataFieldRole,
) -> DataFieldDefinition:
    field = definition.field_for(role)
    if field is None:
        raise BusinessError(
            ErrorCode.DATA_TABLE_INVALID,
            "Resolved data definition is invalid",
            http_status=422,
        )
    return field


def _json_custom_values(
    raw_custom_values: Mapping[str, object],
    definition: ResolvedDataDefinition,
    issues: list[RowIssue],
) -> dict[int, object]:
    """Map JSON field-ID keys onto field IDs, rejecting unknown/system keys."""
    values_by_field_id: dict[int, object] = {}
    for raw_field_id, raw_value in raw_custom_values.items():
        field = _custom_field_for_key(raw_field_id, definition)
        if field is None:
            system_field = _system_field_for_key(raw_field_id, definition)
            if system_field is not None:
                issues.append(
                    RowIssue(
                        "system_field_in_custom_values",
                        "System fields must be supplied as top-level row keys",
                        field_id=system_field.id,
                        column_label=system_field.label,
                    )
                )
            else:
                issues.append(
                    RowIssue(
                        "unknown_custom_field",
                        "custom_values contains an unknown field ID",
                    )
                )
            continue

        if field.system_key != DataFieldRole.CUSTOM:
            issues.append(
                RowIssue(
                    "system_field_in_custom_values",
                    "System fields must be supplied as top-level row keys",
                    field_id=field.id,
                    column_label=field.label,
                )
            )
            continue
        if field.id in values_by_field_id:
            issues.append(
                RowIssue(
                    "duplicate_key",
                    "custom_values contains the same field ID more than once",
                    field_id=field.id,
                    column_label=field.label,
                )
            )
            continue
        values_by_field_id[field.id] = raw_value
    return values_by_field_id


def _convert_custom_values(
    values_by_field_id: Mapping[int, object],
    definition: ResolvedDataDefinition,
    codec: _CellCodec,
    issues: list[RowIssue],
) -> dict[int, FieldValue]:
    """Convert present custom values, enforcing required custom fields."""
    converted: dict[int, FieldValue] = {}
    for field in definition.custom_fields:
        if field.id not in values_by_field_id:
            if field.required:
                issues.append(_required_issue(field))
            continue
        raw_value = values_by_field_id[field.id]
        if raw_value is None:
            if field.required:
                issues.append(_required_issue(field))
            continue
        try:
            converted[field.id] = _convert_custom_value(
                field,
                raw_value,
                definition,
                codec,
            )
        except _ConversionError as error:
            issues.append(_conversion_issue(field, error))
    return converted


def _custom_field_for_key(
    raw_field_id: object,
    definition: ResolvedDataDefinition,
) -> DataFieldDefinition | None:
    if not isinstance(raw_field_id, str) or not _FIELD_ID_PATTERN.fullmatch(
        raw_field_id
    ):
        return None
    field_id = int(raw_field_id)
    if field_id > _INT64_MAX:
        return None
    return definition.fields_by_id.get(field_id)


def _system_field_for_key(
    raw_key: object,
    definition: ResolvedDataDefinition,
) -> DataFieldDefinition | None:
    if not isinstance(raw_key, str):
        return None
    role_by_json_key = {value: role for role, value in _FIXED_JSON_KEYS.items()}
    folded_key = raw_key.casefold()
    for json_key, role in role_by_json_key.items():
        if folded_key == json_key.casefold():
            return definition.field_for(role)
    try:
        role = DataFieldRole(raw_key.upper())
    except ValueError:
        return None
    if role == DataFieldRole.CUSTOM:
        return None
    return definition.field_for(role)


def _convert_custom_value(
    field: DataFieldDefinition,
    raw_value: object,
    definition: ResolvedDataDefinition,
    codec: _CellCodec,
) -> FieldValue:
    if field.field_type == DataFieldType.TEXT:
        return FieldValue(field_id=field.id, text_value=codec.text(raw_value))
    if field.field_type == DataFieldType.INTEGER:
        return FieldValue(
            field_id=field.id,
            integer_value=codec.integer(raw_value),
        )
    if field.field_type == DataFieldType.DECIMAL:
        return FieldValue(
            field_id=field.id,
            decimal_value=codec.decimal(raw_value),
        )
    if field.field_type == DataFieldType.DATE:
        return FieldValue(
            field_id=field.id,
            date_value=codec.date(raw_value),
        )
    if field.field_type == DataFieldType.BOOLEAN:
        return FieldValue(
            field_id=field.id,
            boolean_value=codec.boolean(raw_value),
        )
    if field.field_type == DataFieldType.ENUM:
        return FieldValue(
            field_id=field.id,
            enum_option_id=codec.enum(raw_value, field, definition),
        )
    raise _ConversionError(
        "unsupported_field_type",
        f"Custom field type {field.field_type.value} is not supported",
    )


def _coerce_text(value: object) -> str:
    if not isinstance(value, str):
        raise _ConversionError("invalid_text", "Value must be text")
    if len(value) > _MAX_TEXT_LENGTH:
        raise _ConversionError(
            "text_too_long",
            f"Text must not exceed {_MAX_TEXT_LENGTH} characters",
        )
    return value


def _coerce_integer(value: object) -> int:
    parsed: int | None = None
    if isinstance(value, bool):
        parsed = None
    elif isinstance(value, int):
        parsed = value
    elif isinstance(value, str):
        candidate = value.strip()
        if _INTEGER_PATTERN.fullmatch(candidate):
            digits = candidate.removeprefix("-")
            if len(digits) > 19:
                raise _ConversionError(
                    "integer_out_of_range",
                    "Integer is outside the signed 64-bit range",
                )
            parsed = int(candidate)
    if parsed is None:
        raise _ConversionError("invalid_integer", "Value must be an integer")
    if parsed < _INT64_MIN or parsed > _INT64_MAX:
        raise _ConversionError(
            "integer_out_of_range",
            "Integer is outside the signed 64-bit range",
        )
    return parsed


def _coerce_decimal(value: object) -> Decimal:
    parsed: Decimal
    if isinstance(value, bool) or isinstance(value, float):
        raise _ConversionError("invalid_decimal", "Value must be a decimal")
    if isinstance(value, Decimal):
        parsed = value
    elif isinstance(value, int):
        parsed = Decimal(value)
    elif isinstance(value, str):
        candidate = value.strip()
        if not _DECIMAL_PATTERN.fullmatch(candidate):
            raise _ConversionError("invalid_decimal", "Value must be a decimal")
        try:
            parsed = Decimal(candidate)
        except InvalidOperation as error:
            raise _ConversionError(
                "invalid_decimal",
                "Value must be a decimal",
            ) from error
    else:
        raise _ConversionError("invalid_decimal", "Value must be a decimal")

    if not parsed.is_finite():
        raise _ConversionError("invalid_decimal", "Value must be a finite decimal")
    exponent = parsed.as_tuple().exponent
    if not isinstance(exponent, int):
        raise _ConversionError("invalid_decimal", "Value must be a finite decimal")
    if exponent < -_MAX_DECIMAL_PLACES:
        raise _ConversionError(
            "too_many_decimal_places",
            f"Decimal must not exceed {_MAX_DECIMAL_PLACES} decimal places",
        )
    if parsed != 0 and parsed.adjusted() >= _MAX_DECIMAL_INTEGER_DIGITS:
        raise _ConversionError(
            "decimal_out_of_range",
            "Decimal is outside the supported DECIMAL(30,8) range",
        )
    return parsed


def _coerce_date(value: object) -> date:
    if isinstance(value, datetime) or not isinstance(value, (date, str)):
        raise _ConversionError("invalid_date", "Value must be an ISO date")
    if isinstance(value, date):
        return value
    candidate = value.strip()
    if not _ISO_DATE_PATTERN.fullmatch(candidate):
        raise _ConversionError("invalid_date", "Value must be an ISO date")
    try:
        return date.fromisoformat(candidate)
    except ValueError as error:
        raise _ConversionError(
            "invalid_date",
            "Value must be a valid ISO date",
        ) from error


def _coerce_boolean(value: object) -> bool:
    if not isinstance(value, bool):
        raise _ConversionError("invalid_boolean", "Value must be a boolean")
    return value


def _coerce_enum(
    value: object,
    field: DataFieldDefinition,
    definition: ResolvedDataDefinition,
) -> int:
    if not isinstance(value, str):
        raise _ConversionError("invalid_enum", "Value must be an enum label")
    label = value.strip()
    for option in definition.options_for(field.id):
        if option.label == label:
            return option.id
    raise _ConversionError("invalid_enum", "Value is not an allowed enum label")


def _coerce_customer_id(value: object) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise _ConversionError("invalid_customer", "Customer ID is invalid")
    return value


def _resolve_customer(
    customer_id: int,
    customers_by_id: Mapping[int, CustomerIdentity],
    allowed_customer_ids: set[int] | None,
) -> CustomerIdentity:
    identity = customers_by_id.get(customer_id)
    if identity is None or identity.status != ContactStatus.ACTIVE:
        raise _ConversionError("invalid_customer", "Customer is unavailable")
    if allowed_customer_ids is not None and customer_id not in allowed_customer_ids:
        raise _ConversionError(
            "customer_out_of_scope",
            "Customer is outside your data scope",
        )
    return identity


def _required_issue(field: DataFieldDefinition) -> RowIssue:
    return RowIssue(
        "required",
        f"Field {field.label} is required",
        field_id=field.id,
        column_label=field.label,
    )


def _conversion_issue(
    field: DataFieldDefinition,
    error: _ConversionError,
) -> RowIssue:
    return RowIssue(
        error.error_code,
        f"Field {field.label}: {error.message}",
        field_id=field.id,
        column_label=field.label,
    )


def _duplicate_issue_rows(
    raw: Mapping[str, object],
    definition: ResolvedDataDefinition,
    row_number: int,
) -> tuple[RowIssue, ...]:
    issues = _duplicate_object_issues(raw, definition)
    custom_values = raw.get("custom_values")
    if isinstance(custom_values, Mapping):
        issues.extend(_duplicate_object_issues(custom_values, definition))
    return tuple(
        RowIssue(
            error_code=issue.error_code,
            message=issue.message,
            field_id=issue.field_id,
            column_label=issue.column_label,
            sheet_name=issue.sheet_name,
            row_number=row_number,
        )
        for issue in issues
    )


def _duplicate_object_issues(
    raw: Mapping[str, object],
    definition: ResolvedDataDefinition,
) -> list[RowIssue]:
    if not isinstance(raw, _JsonObject):
        return []
    issues: list[RowIssue] = []
    for key in raw.duplicate_keys:
        field = _field_for_duplicate_key(key, definition)
        issues.append(
            RowIssue(
                "duplicate_key",
                "Object contains the same key more than once",
                field_id=field.id if field is not None else None,
                column_label=field.label if field is not None else None,
            )
        )
    return issues


def _field_for_duplicate_key(
    key: str,
    definition: ResolvedDataDefinition,
) -> DataFieldDefinition | None:
    role_by_json_key = {value: role for role, value in _FIXED_JSON_KEYS.items()}
    if key in role_by_json_key:
        return definition.field_for(role_by_json_key[key])
    return _custom_field_for_key(key, definition)


# ---------------------------------------------------------------------------
# Excel transport
# ---------------------------------------------------------------------------


class _ExcelCellCodec:
    """Normalize Excel-native cell forms into the shared domain values.

    openpyxl decodes a cell as ``str``, ``bool``, ``int``, ``float``,
    ``datetime``/``date``/``time``, or ``None``. Only syntax a spreadsheet can
    genuinely carry is normalized here; every acceptance rule (text length,
    int64 and DECIMAL(30,8) ranges, date legality, enum labels, customer scope)
    still runs in the shared conversion functions.
    """

    def text(self, value: object) -> str:
        return _coerce_text(value)

    def integer(self, value: object) -> int:
        if isinstance(value, bool):
            raise _ConversionError("invalid_integer", "Value must be an integer")
        # Excel stores every number as an IEEE-754 double, so integral 42.0 is
        # the same cell value as 42; a fractional value is not an integer.
        if isinstance(value, float):
            if not value.is_integer():
                raise _ConversionError(
                    "invalid_integer",
                    "Value must be an integer",
                )
            return _coerce_integer(int(value))
        return _coerce_integer(value)

    def decimal(self, value: object) -> Decimal:
        if isinstance(value, float):
            if value != value or value in (_INFINITY, -_INFINITY):
                raise _ConversionError(
                    "invalid_decimal",
                    "Value must be a finite decimal",
                )
            # Spreadsheet numbers are binary doubles. ``repr`` gives the
            # shortest round-trip decimal, so 1942.51 becomes exactly
            # Decimal("1942.51") instead of a 50-digit binary expansion.
            return _coerce_decimal(repr(value))
        return _coerce_decimal(value)

    def date(self, value: object) -> date:
        if isinstance(value, datetime):
            if value.time() != _MIDNIGHT:
                raise _ConversionError(
                    "invalid_date",
                    "Value must be a date without a time",
                )
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            candidate = value.strip()
            if _COMPACT_DATE_PATTERN.fullmatch(candidate):
                return _parse_compact_date(candidate)
            return _coerce_date(candidate)
        if isinstance(value, int) and not isinstance(value, bool):
            return _parse_compact_date(str(value))
        return _coerce_date(value)

    def boolean(self, value: object) -> bool:
        if isinstance(value, bool):
            return value
        if isinstance(value, str):
            candidate = value.strip().casefold()
            if candidate in _EXCEL_BOOLEAN_TRUE:
                return True
            if candidate in _EXCEL_BOOLEAN_FALSE:
                return False
            raise _ConversionError("invalid_boolean", "Value must be a boolean")
        if isinstance(value, int) and value in (0, 1):
            return value == 1
        raise _ConversionError("invalid_boolean", "Value must be a boolean")

    def enum(
        self,
        value: object,
        field: DataFieldDefinition,
        definition: ResolvedDataDefinition,
    ) -> int:
        return _coerce_enum(value, field, definition)

    def customer_id(
        self,
        value: object,
        customers_by_name: Mapping[str, CustomerIdentity],
    ) -> int:
        if not isinstance(value, str):
            raise _ConversionError("invalid_customer", "Customer is unavailable")
        identity = customers_by_name.get(customer_name_lookup_key(value))
        if identity is None:
            raise _ConversionError("invalid_customer", "Customer is unavailable")
        return identity.id


_EXCEL_CELL_CODEC = _ExcelCellCodec()
_MIDNIGHT = datetime.min.time()
_INFINITY = float("inf")
_COMPACT_DATE_PATTERN = re.compile(r"^[0-9]{8}$")
_HEADER_ROW_INDEX = 1


def _parse_compact_date(candidate: str) -> date:
    if not _COMPACT_DATE_PATTERN.fullmatch(candidate):
        raise _ConversionError("invalid_date", "Value must be a valid date")
    try:
        return date(int(candidate[0:4]), int(candidate[4:6]), int(candidate[6:8]))
    except ValueError as error:
        raise _ConversionError(
            "invalid_date",
            "Value must be a valid date",
        ) from error


def _excel_error(message: str) -> BusinessError:
    return BusinessError(
        ErrorCode.DATA_IMPORT_INVALID,
        message,
        http_status=422,
    )


def _excel_invalid() -> BusinessError:
    return _excel_error(_EXCEL_INVALID_MESSAGE)


def _read_excel_workbook(path: Path) -> object:
    """Validate the XLSX container, then return a read-only workbook.

    Rejection happens before any cell is converted: size, Zip signature, ZIP
    container, workbook part, macro part/content type, and later worksheet
    counts. The workbook is opened with ``read_only=True`` and
    ``data_only=False`` so formulas stay visible as formula cells and cached
    results are never read.
    """
    try:
        if not path.is_file():
            raise _excel_invalid()
        size = path.stat().st_size
    except OSError as error:
        raise _excel_invalid() from error
    if size <= 0 or size > _XLSX_MAX_BYTES:
        raise _excel_invalid()

    try:
        with path.open("rb") as source:
            signature = source.read(len(_XLSX_SIGNATURE))
    except OSError as error:
        raise _excel_invalid() from error
    if signature != _XLSX_SIGNATURE:
        raise _excel_invalid()

    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if _WORKBOOK_PART not in names or _MACRO_PART in names:
                raise _excel_invalid()
            if _MACRO_CONTENT_TYPE in archive.read("[Content_Types].xml"):
                raise _excel_invalid()
            if any(name.startswith(_CHARTSHEET_PREFIX) for name in names):
                raise _excel_error(_EXCEL_SHEET_COUNT_MESSAGE)
            if _WORKBOOK_RELS_PART in names:
                _require_no_chartsheet(archive.read(_WORKBOOK_RELS_PART))
            _require_single_visible_sheet(archive.read(_WORKBOOK_PART))
    except (
        ElementTree.ParseError,
        KeyError,
        OSError,
        RuntimeError,
        ValueError,
        zipfile.BadZipFile,
    ) as error:
        raise _excel_invalid() from error

    try:
        return cast(
            object,
            load_workbook(
                path,
                read_only=True,
                data_only=False,
                keep_links=False,
            ),
        )
    except (
        ElementTree.ParseError,
        InvalidFileException,
        KeyError,
        OSError,
        RuntimeError,
        TypeError,
        ValueError,
    ) as error:
        raise _excel_invalid() from error


def _require_no_chartsheet(workbook_rels_xml: bytes) -> None:
    """Reject chartsheets by relationship type, not by part path.

    openpyxl 3.1 cannot attach chartsheet drawings in read-only mode, so a
    chartsheet must be refused before ``load_workbook`` runs. Checking the
    relationship type catches chartsheets stored outside ``xl/chartsheets/``.
    """
    root = ElementTree.fromstring(workbook_rels_xml)
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] != "Relationship":
            continue
        if element.get("Type", "").endswith("/chartsheet"):
            raise _excel_error(_EXCEL_SHEET_COUNT_MESSAGE)


def _require_single_visible_sheet(workbook_xml: bytes) -> None:
    """Count declared worksheets from raw ``workbook.xml`` before openpyxl.

    Chartsheets are rejected first (openpyxl 3.1 crashes while attaching their
    drawings), and the declared sheet list must be exactly one visible
    worksheet.
    """
    root = ElementTree.fromstring(workbook_xml)
    sheets = [
        element for element in root.iter() if element.tag.rsplit("}", 1)[-1] == "sheet"
    ]
    if len(sheets) != 1:
        raise _excel_error(_EXCEL_SHEET_COUNT_MESSAGE)
    state = sheets[0].get("state")
    if state not in (None, "visible"):
        raise _excel_error(_EXCEL_SHEET_COUNT_MESSAGE)


def _import_worksheet(workbook: object) -> object:
    """Return the single visible worksheet or reject the whole workbook."""
    worksheets = getattr(workbook, "worksheets", None)
    chartsheets = getattr(workbook, "chartsheets", None)
    if (
        not isinstance(worksheets, list)
        or not isinstance(chartsheets, list)
        or chartsheets
        or len(worksheets) != 1
        or getattr(worksheets[0], "sheet_state", None) != "visible"
    ):
        raise _excel_error(_EXCEL_SHEET_COUNT_MESSAGE)
    return worksheets[0]


def _cell_value(cell: object) -> object:
    """Return a raw cell value, rejecting formula and error cells.

    Rejection is driven only by the cell's stored type, never by its text:
    a formula cell is ``f`` and an error cell is ``e``. Text that merely looks
    like a formula (for example a legitimate ``"=1+1"`` identifier typed into a
    text field) remains ordinary text, so TEXT fields keep every character.
    Formulas are refused before any value is read, which is what makes cached
    results unusable.
    """
    data_type = getattr(cell, "data_type", None)
    if data_type == "f":
        raise _excel_error(_EXCEL_FORMULA_MESSAGE)
    if data_type == "e":
        raise _excel_error(_EXCEL_ERROR_VALUE_MESSAGE)
    return getattr(cell, "value", None)


def _row_values(row: tuple[object, ...], column_count: int) -> tuple[object, ...]:
    """Reject values past the configured columns, then pad missing cells."""
    values: list[object] = []
    for index, cell in enumerate(row):
        value = _cell_value(cell)
        if index >= column_count:
            if value is not None:
                raise _excel_error(_EXCEL_EXTRA_VALUE_MESSAGE)
            continue
        values.append(value)
    values.extend([None] * (column_count - len(values)))
    return tuple(values)


def _worksheet_rows(worksheet: object) -> Iterator[tuple[int, tuple[object, ...]]]:
    """Yield ``(worksheet row number, cells)`` without trusting dimensions.

    ``reset_dimensions`` makes read-only iteration stream the real ``sheetData``
    rows, so trailing styled-but-empty rows never inflate the row count.
    """
    reset = getattr(worksheet, "reset_dimensions", None)
    if callable(reset):
        reset()
    iterator = getattr(worksheet, "iter_rows", None)
    if not callable(iterator):
        raise _excel_invalid()
    try:
        yield from enumerate(
            cast(Iterator[tuple[object, ...]], iterator(min_row=1, min_col=1)),
            start=1,
        )
    except (
        ElementTree.ParseError,
        KeyError,
        OSError,
        RuntimeError,
        TypeError,
        ValueError,
    ) as error:
        raise _excel_invalid() from error


def _row_at(
    worksheet: object,
    target_row_number: int,
) -> tuple[object, ...] | None:
    """Read one worksheet row, closing the streaming reader afterwards.

    A read-only worksheet opens a ZIP member per iteration, so an abandoned
    half-consumed iterator must be closed explicitly instead of waiting for
    garbage collection.
    """
    rows = _worksheet_rows(worksheet)
    try:
        for row_number, row in rows:
            if row_number == target_row_number:
                return row
            if row_number > target_row_number:
                return None
        return None
    finally:
        close = getattr(rows, "close", None)
        if callable(close):
            close()


def _header_columns(
    worksheet: object,
    definition: ResolvedDataDefinition,
) -> tuple[int, dict[int, int]]:
    """Map field IDs to column indexes from the first row of the sheet.

    Matching uses the project's label normalization (trim only), is
    order-independent, and rejects missing, duplicate, unknown, blank, and
    non-text headers.
    """
    fields_by_label = {
        _normalized_label(field.label): field for field in definition.fields
    }
    header_row = _row_at(worksheet, _HEADER_ROW_INDEX)
    if header_row is None or not any(
        _cell_value(cell) is not None for cell in header_row
    ):
        raise _excel_error(_EXCEL_HEADER_MESSAGE)

    values = [_cell_value(cell) for cell in header_row]
    last_label_index = max(
        (index for index, value in enumerate(values) if value is not None),
        default=-1,
    )
    if last_label_index < 0:
        raise _excel_error(_EXCEL_HEADER_MESSAGE)

    columns: dict[int, int] = {}
    seen: set[str] = set()
    for column_index in range(last_label_index + 1):
        value = values[column_index]
        if value is None:
            # An interior blank column would silently swallow any value below
            # it, so a gap in the label block is treated as a missing header.
            raise _excel_error(_EXCEL_HEADER_MESSAGE)
        if not isinstance(value, str):
            raise _excel_error(_EXCEL_HEADER_MESSAGE)
        label = _normalized_label(value)
        if not label or label in seen:
            raise _excel_error(_EXCEL_HEADER_MESSAGE)
        seen.add(label)
        field = fields_by_label.get(label)
        if field is None:
            raise _excel_error(_EXCEL_HEADER_MESSAGE)
        columns[field.id] = column_index

    if len(columns) != len(fields_by_label):
        raise _excel_error(_EXCEL_HEADER_MESSAGE)
    # Width stops at the last configured label: anything a data row writes past
    # it is a stray value, not silently ignored padding.
    return last_label_index + 1, columns


def _normalized_label(label: str) -> str:
    """The project's label semantics: trim, exact case, exact interior."""
    return label.strip()


def iter_excel_rows(
    path: Path,
    definition: ResolvedDataDefinition,
    customer_index: Mapping[str, CustomerIdentity],
    allowed_customer_ids: set[int] | None,
) -> Iterator[ParsedImportRow]:
    """Yield one validated or issue-carrying result per non-empty Excel row.

    Headers come from row one, columns may be in any order, and every
    configured label must appear exactly once. Empty rows are skipped, and each
    result carries its true worksheet row number. Row issues name the sheet,
    row, configured column label, and field ID, and never repeat a raw value.

    Structural problems (invalid container, wrong sheet count, invalid headers,
    formulas, error values, values past the configured columns, or more than
    ``_MAX_ROWS`` data rows) raise one ``BusinessError(9006)`` for the file.
    Per-row conversion problems are returned as ``ParsedImportRow.issues``.

    The function is synchronous and builds the customer-ID index exactly once;
    Task 9 is responsible for running it in a worker thread.
    """
    workbook = _read_excel_workbook(path)
    try:
        worksheet = _import_worksheet(workbook)
        sheet_name = getattr(worksheet, "title", None)
        if not isinstance(sheet_name, str):
            sheet_name = None
        column_count, columns = _header_columns(worksheet, definition)
        customers_by_id = _customer_index_by_id(customer_index)

        data_rows = 0
        rows = _worksheet_rows(worksheet)
        try:
            for row_number, row in rows:
                if row_number <= _HEADER_ROW_INDEX:
                    continue
                values = _row_values(row, column_count)
                if all(value is None for value in values):
                    continue
                data_rows += 1
                if data_rows > _MAX_ROWS:
                    raise _excel_error(
                        f"Excel file must not exceed {_MAX_ROWS} data rows",
                    )
                yield _parse_excel_row(
                    values,
                    row_number,
                    sheet_name,
                    definition,
                    columns,
                    customer_index,
                    customers_by_id,
                    allowed_customer_ids,
                )
        finally:
            # A caller may stop consuming early (for example after the first
            # error); release the read-only ZIP member deterministically.
            close_rows = getattr(rows, "close", None)
            if callable(close_rows):
                close_rows()
    finally:
        close = getattr(workbook, "close", None)
        if callable(close):
            close()


def _parse_excel_row(
    values: tuple[object, ...],
    row_number: int,
    sheet_name: str | None,
    definition: ResolvedDataDefinition,
    columns: Mapping[int, int],
    customers_by_name: Mapping[str, CustomerIdentity],
    customers_by_id: Mapping[int, CustomerIdentity],
    allowed_customer_ids: set[int] | None,
) -> ParsedImportRow:
    system_raw: dict[DataFieldRole, object] = {}
    custom_raw: dict[int, object] = {}
    for field_id, column_index in columns.items():
        field = definition.fields_by_id[field_id]
        value = values[column_index] if column_index < len(values) else None
        if field.system_key == DataFieldRole.CUSTOM:
            custom_raw[field.id] = value
        else:
            system_raw[field.system_key] = value

    issues: list[RowIssue] = []
    result = _validate_values(
        system_raw,
        custom_raw,
        definition,
        customers_by_id,
        _EXCEL_CELL_CODEC,
        issues,
        allowed_customer_ids=allowed_customer_ids,
        customers_by_name=customers_by_name,
    )
    if isinstance(result, ValidatedImportRow):
        return ParsedImportRow(row_number=row_number, row=result)
    return ParsedImportRow(
        row_number=row_number,
        issues=tuple(
            RowIssue(
                error_code=issue.error_code,
                message=issue.message,
                field_id=issue.field_id,
                column_label=issue.column_label,
                sheet_name=sheet_name,
                row_number=row_number,
            )
            for issue in result
        ),
    )


def build_excel_template(definition: ResolvedDataDefinition) -> bytes:
    """Return XLSX bytes with one visible sheet and one configured header row.

    The template contains exactly ``definition.labels`` in field order and no
    example data, so an operator fills data from row two and submits the file
    unchanged.
    """
    workbook = Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = "Data"
        worksheet.append(list(definition.labels))
        for index in range(1, len(definition.labels) + 1):
            worksheet.cell(row=1, column=index).font = Font(bold=True)
        worksheet.freeze_panes = "A2"

        buffer = BytesIO()
        workbook.save(buffer)
    finally:
        workbook.close()
    return buffer.getvalue()


__all__: tuple[str, ...] = (
    "CustomerIdentity",
    "FieldValue",
    "ParsedImportRow",
    "RowIssue",
    "ValidatedImportRow",
    "build_excel_template",
    "iter_excel_rows",
    "iter_json_rows",
    "json_import_row_count",
    "load_customer_name_index",
    "validate_raw_row",
)
